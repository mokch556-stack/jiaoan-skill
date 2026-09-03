# -*- coding: utf-8 -*-
"""dump 教案全文：所有段落 + 所有表格到文件"""
import docx, sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
path = r'E:\qwenpaw生成\教案\人工智能通识\人工智能通识_第1次课教案.docx'
out = r'E:\qwenpaw\workspaces\default\temp\教案v4.12_full_dump.txt'
d = docx.Document(path)
lines = []
lines.append('===== 段落 =====')
for i, p in enumerate(d.paragraphs):
    t = p.text
    if t.strip():
        lines.append(f'P{i}: {t}')
lines.append('')
lines.append(f'===== 表格 {len(d.tables)} 个 =====')
for ti, tb in enumerate(d.tables):
    lines.append(f'--- 表{ti} ---')
    seen = set()
    for ri, row in enumerate(tb.rows):
        for ci, c in enumerate(row.cells):
            if id(c._tc) in seen: continue
            seen.add(id(c._tc))
            txt = c.text.strip()
            if txt:
                lines.append(f'[T{ti}R{ri}C{ci}] {txt}')
    lines.append('')
with open(out, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))
print('已输出', out, '共', len(lines), '行')
